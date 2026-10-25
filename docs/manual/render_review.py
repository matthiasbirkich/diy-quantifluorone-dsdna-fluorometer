"""Create portable Word editions directly from the authoritative Quarto chapters."""
from pathlib import Path
import re
import argparse
from urllib.parse import quote, unquote
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent.parent
OUTPUT = ROOT / 'review'
OUTPUT.mkdir(exist_ok=True)
parser=argparse.ArgumentParser()
parser.add_argument('--book',action='store_true')
parser.add_argument('--chapters',nargs='+',type=int,default=[4,5])
args=parser.parse_args()
if args.book:args.chapters=list(range(1,15))
chapter_names={1:'01_instrument_overview.qmd',2:'02_safety.qmd',3:'03_measurement_principles.qmd',4:'04_components_and_files.qmd',5:'05_assembly.qmd',6:'06_software_installation.qmd',7:'07_startup_and_controls.qmd',8:'08_calibration.qmd',9:'09_routine_analysis.qmd',10:'10_results_and_qc.qmd',11:'11_performance.qmd',12:'12_maintenance_and_troubleshooting.qmd',13:'13_sybr_safe.qmd',14:'14_references.qmd'}
selected=[(n,chapter_names[n]) for n in args.chapters]
subtitle=('Assembly and Operating Manual' if args.book else 'SYBR Safe and References' if args.chapters==[13,14] else 'Instrument Overview Safety and Measurement Principles' if args.chapters==[1,2,3] else 'Analysis Performance and Maintenance' if args.chapters==[8,9,10,11,12] else ('Calibration Routine Analysis and Quality Control' if args.chapters==[8,9,10] else ('Software Installation and Instrument Operation' if args.chapters==[6,7] else 'Components and Assembly')))
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    style=doc.styles[name]
    style.font.name='Calibri'
    style.font.color.rgb=RGBColor(0,0,0)
    if style._element.pPr is not None:
        for e in list(style._element.pPr):
            if e.tag==qn('w:pBdr'):style._element.pPr.remove(e)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(4 if min(args.chapters)>=8 else 6)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Subtitle'].font.size=Pt(12)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Caption'].font.size=Pt(9)
doc.add_paragraph('DIY QuantiFluorONE dsDNA Fluorometer','Title')
doc.add_paragraph(subtitle,'Subtitle')
doc.add_paragraph('Dipl.-Ing. Matthias Birkicht')
doc.add_paragraph('Centre for Tropical Marine Research')
if args.book:
    doc.add_paragraph('24 October 2026')
    doc.add_heading('About this manual',level=2)



def inline(p,text,source):
    pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*.*?\*\*|`[^`]+`)'
    for token in re.split(pattern,text):
        if not token:continue
        link=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
        if link:
            label,target=link.groups()
            if not target.startswith(('http:','https:','#')):
                rel=(source.parent/unquote(target)).resolve().relative_to(REPO).as_posix()
                target='https://github.com/matthiasbirkich/diy-quantifluorone-dsdna-fluorometer/blob/main/'+quote(rel,safe='/')
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'185A82');pr.append(color);r.append(pr)
            t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
        elif token.startswith('**'):
            p.add_run(token[2:-2]).bold=True
        elif token.startswith('`'):
            run=p.add_run(token[1:-1]);run.font.name='Consolas';run.font.size=Pt(9)
        else:p.add_run(token)

if args.book:
    intro_source=ROOT/'index.qmd'
    for paragraph in intro_source.read_text().split('\n\n'):
        if paragraph.startswith('#'):continue
        if paragraph.strip():inline(doc.add_paragraph(),paragraph.strip(),intro_source)

for number,name in selected:
    source=ROOT/'chapters'/name
    content=re.sub(r'\A---\n.*?\n---\n','',source.read_text(),flags=re.S)
    lines=content.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        heading=re.match(r'^(#{1,3}) (.*?)(?: \{#[^}]+\})?$',line)
        image=re.match(r'^!\[([^\]]*)\]\(([^)]+)\)',line)
        if heading:
            level=len(heading[1]);title=heading[2]
            if level==1:title=f'{number} {title}'
            p=doc.add_heading(title,level=level)
            if level==1 and number!=selected[0][0]:p.paragraph_format.page_break_before=True
            i+=1
        elif image:
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str((source.parent/image[2]).resolve()),width=Cm(8 if 'assembled_instrument_front' in image[2] else (9 if 'cad' in image[2] else (8 if 'detection_side' in image[2] else 12))))
            cap=doc.add_paragraph(style='Caption');cap.add_run(image[1]);i+=1
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',v) for v in cells):rows.append(cells)
                i+=1
            table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Table Grid';table.autofit=False
            widths={3:[2.1,4.8,10.1],4:[4.2,3.2,3.2,6.4]}.get(len(rows[0]),[17/len(rows[0])]*len(rows[0]))
            if rows[0][0]=='Task' and len(rows[0])==2:widths=[12,5]
            if rows[0][0].startswith('Quantity for') and len(rows[0])==2:widths=[9,8]
            if rows[0][0]=='Component' and len(rows[0])==3:widths=[3.5,7,6.5]
            if rows[0][0]=='Action' and len(rows[0])==3:widths=[4,6,7]
            if rows[0][0] in ['Route','Tube','Measurement design','Observation','Active route','Measure'] and len(rows[0])==3:widths=[5,6,6]
            for j,w in enumerate(widths):table.columns[j].width=Cm(w)
            for rno,row in enumerate(rows):
                cells=table.rows[0].cells if rno==0 else table.add_row().cells
                table.rows[rno]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for j,value in enumerate(row):
                    cells[j].width=Cm(widths[j])
                    tcpr=cells[j]._tc.get_or_add_tcPr()
                    borders=OxmlElement('w:tcBorders')
                    for side in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    tcpr.append(borders)
                    if rno==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EFF5');tcpr.append(shade)
                    p=cells[j].paragraphs[0];p.paragraph_format.space_after=Pt(3);inline(p,value,source)
                    for run in p.runs:run.font.size=Pt(9);run.bold=rno==0
                if rno==0:
                    pr=table.rows[0]._tr.get_or_add_trPr();pr.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph()
        elif re.match(r'^(\d+\. |- \[ \] |- )',line):
            p=doc.add_paragraph();inline(p,line.replace('- [ ] ','☐ ',1) if line.startswith('- [ ]') else line,source);i+=1
        else:
            values=[line];i+=1
            while i<len(lines) and lines[i].strip() and not re.match(r'^(#|!\[|\||\d+\. |- )',lines[i].strip()):values.append(lines[i].strip());i+=1
            inline(doc.add_paragraph(),' '.join(values),source)
path=OUTPUT/('manual_source.docx' if args.book else 'Build_Manual_Chapters_'+'_'.join(map(str,args.chapters))+'.docx')
doc.save(path)
print(path)
