"""Polish a Word edition from Quarto sources and populate chapter contents from PDF pagination.

Run after render_review.py --book. Render once, then pass --pagination-pdf and
render again. All chapter wording remains in the QMD sources.
"""
from pathlib import Path
from urllib.parse import unquote,quote
import re,argparse
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
REPO=ROOT.parent.parent
p=argparse.ArgumentParser();p.add_argument('--input',type=Path,default=ROOT/'review/manual_source.docx');p.add_argument('--output',type=Path,default=ROOT/'review/DIY_QuantiFluorONE_Assembly_and_Operating_Manual.docx');p.add_argument('--pagination-pdf',type=Path);a=p.parse_args()
d=Document(a.input)
# Replace relative repository URLs in the portable Word edition.
for rel in d.part.rels.values():
 if rel.reltype==RT.HYPERLINK and not re.match(r'^[a-z]+:',rel.target_ref) and not rel.target_ref.startswith('#'):
  target=(ROOT/'chapters'/unquote(rel.target_ref)).resolve()
  if not target.exists():raise ValueError('Missing linked source '+str(target))
  rel._target='https://github.com/matthiasbirkich/diy-quantifluorone-dsdna-fluorometer/blob/main/'+quote(target.relative_to(REPO).as_posix(),safe='/')
# Set a consistent handbook style.
for sec in d.sections:
 sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
 sec.footer_distance=Cm(.8)
 fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 run=fp.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
for name in ['Normal','Body Text','First Paragraph','Compact']:
 if name in d.styles:
  s=d.styles[name];s.font.name='Calibri';s.font.size=Pt(10);s.paragraph_format.space_after=Pt(4);s.paragraph_format.line_spacing=1.0
for name,size in [('Title',25),('Subtitle',15),('Heading 1',17),('Heading 2',12),('Heading 3',11),('Caption',9),('TOC Heading',17)]:
 if name in d.styles:
  s=d.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
  s.font.underline=False
  for attr in ('color','u'):
   e=s.element.find('w:rPr/w:'+attr, namespaces=s.element.nsmap)
   if e is not None and attr=='color':e.attrib.pop(qn('w:themeColor'),None)
for name in ['Heading 1','Heading 2','Heading 3']:
 d.styles[name].paragraph_format.space_before=Pt(10);d.styles[name].paragraph_format.space_after=Pt(5);d.styles[name].paragraph_format.keep_with_next=True
# Avoid a long caption or image pushing a large block to the next page.
for shape in d.inline_shapes:
 w,h=shape.width,shape.height
 limitw,limith=Cm(12),Cm(9.5)
 if h>limith or w>limitw:
  factor=min(limitw/w,limith/h);shape.width=int(w*factor);shape.height=int(h*factor)
for para in d.paragraphs:
 if para._p.xpath('.//w:drawing'):
  para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
  nxt=para._p.getnext()
  if nxt is not None and nxt.tag==qn('w:p'):
   from docx.text.paragraph import Paragraph
   caption=Paragraph(nxt,d._body);caption.style='Caption';caption.paragraph_format.keep_with_next=False
 for r in para.runs:
  if r.font.color is not None and para.style.name.startswith(('Title','Subtitle','Heading')):r.font.color.rgb=RGBColor(0,0,0)
# Table widths follow the content schema, never a numeric first column by default.
for table in d.tables:
 headers=[c.text.strip() for c in table.rows[0].cells];n=len(headers);first=headers[0]
 widths=[17/n]*n
 if n==2:widths=[7,10]
 if first=='Task' and n==2:widths=[11.5,5.5]
 if first=='Quantity for the high standard':widths=[10,7]
 if first=='Quantity':widths=[2.2,5,9.8]
 if first=='Part':widths=[4.3,3.3,3.2,6.2]
 if first=='Stage':widths=[1.4,6,9.6]
 if first=='Device':widths=[6,3,8]
 if first=='Component' and n==3:widths=[4,6.3,6.7]
 if first=='Action' and n==3:widths=[4.3,5.8,6.9]
 if first=='Control':widths=[2.3,7.3,7.4]
 if first=='Entry':widths=[3.4,6.4,7.2]
 if first in ['Route','Active route','Measurement design','Observation'] and n==3:widths=[5,5.8,6.2]
 if first=='Tube' and n==3:widths=[4.5,6,6.5]
 if first=='Tube' and n==4:widths=[3.4,3.1,4.7,5.8]
 if headers==['Tube','Master mix','Added solution','Purpose']:widths=[3.2,2.2,5.6,6]
 if first=='Concentration basis':widths=[7,5,5]
 if first=='Setting in the assay object':widths=[8,9]
 table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
 for col,width in zip(table.columns,widths):col.width=Cm(width)
 for ri,row in enumerate(table.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,cell in enumerate(row.cells):
   cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr()
   for tag in ['tcBorders','tcMar','shd']:
    for e in list(tcpr.findall(qn('w:'+tag))):tcpr.remove(e)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcpr.append(borders)
   mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'70' if side in ['top','bottom'] else '90');e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
   if ri==0:
    e=OxmlElement('w:shd');e.set(qn('w:fill'),'E7EFF5');tcpr.append(e)
   for para in cell.paragraphs:
    para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.0;para.paragraph_format.keep_with_next=False
    for run in para.runs:run.font.size=Pt(9);run.bold=(ri==0)
# Cover includes the preface. Contents follows on its own page.
body=d._element.body
for sdt in list(body.findall(qn('w:sdt'))):body.remove(sdt)
chapters=[]
for para in d.paragraphs:
 if para.style.name=='Heading 1':
  if para.text=='About this manual':
   para.style='Heading 2';para.paragraph_format.space_before=Pt(20)
  else:
   match=re.match(r'^(\d+)\.?\s+(.*)',para.text)
   if not match:raise ValueError(para.text)
   num=int(match[1]);anchor='chapter_'+str(num)
   b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(9000+num));b.set(qn('w:name'),anchor);para._p.insert(0,b)
   e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(9000+num));para._p.append(e)
   para.paragraph_format.page_break_before=True
   chapters.append((num,para.text,anchor,para))
if len(chapters)!=14:raise ValueError('Expected fourteen chapters')
page_map={}
if a.pagination_pdf:
 import fitz
 pdf=fitz.open(a.pagination_pdf)
 for num,title,anchor,_ in chapters:
  needle=re.sub(r'\s+',' ',title).strip()
  matches=[i+1 for i,page in enumerate(pdf) if needle in re.sub(r'\s+',' ',page.get_text())]
  if not matches:raise ValueError('Cannot locate '+title)
  page_map[num]=matches[-1] # The first match can be the contents entry.
# Number second-level headings consistently in the assembled edition.
chapter_num=0;subnum=0
for para in d.paragraphs:
    if para.style.name=='Heading 1':
        chapter_num=int(re.match(r'^(\d+)',para.text)[1]);subnum=0
    elif para.style.name=='Heading 2' and chapter_num:
        subnum+=1
        if not re.match(r'^\d+\.\d+ ',para.text):para.text=f'{chapter_num}.{subnum} '+para.text
# Build a deterministic linked contents page from the chapter headings.
sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');obj=OxmlElement('w:docPartObj');gal=OxmlElement('w:docPartGallery');gal.set(qn('w:val'),'Table of Contents');obj.append(gal);pr.append(obj);sdt.append(pr)
content=OxmlElement('w:sdtContent');sdt.append(content)
from docx.text.paragraph import Paragraph
h=OxmlElement('w:p');content.append(h);hp=Paragraph(h,d._body);hp.style='TOC Heading';hp.text='Contents';hp.paragraph_format.page_break_before=True;hp.paragraph_format.space_after=Pt(16)
def char(p,kind):
 r=OxmlElement('w:r');e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),kind);r.append(e);p.append(r)
for i,(num,title,anchor,_) in enumerate(chapters):
 el=OxmlElement('w:p');content.append(el);p=Paragraph(el,d._body);p.style='Normal';p.paragraph_format.space_after=Pt(10);p.paragraph_format.tab_stops.add_tab_stop(Cm(16.7),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),anchor);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=title;r.append(t);link.append(r);el.append(link)
 p.add_run('\t'+str(page_map.get(num,0)))
chapters[0][3]._p.addprevious(sdt)
# Keep the source and acknowledgement entries compact without changing type size.
in_references=False
for para in d.paragraphs:
 if para.style.name=='Heading 1':in_references=para.text.startswith('14 ')
 if in_references:
  para.paragraph_format.space_after=Pt(2)
  if para.style.name.startswith('Heading'):para.paragraph_format.space_before=Pt(8)
a.output.parent.mkdir(parents=True,exist_ok=True);d.core_properties.title='DIY QuantiFluorONE dsDNA Fluorometer Assembly and Operating Manual';d.core_properties.author='Dipl.-Ing. Matthias Birkicht';d.core_properties.subject='Construction calibration and routine dsDNA measurement';d.save(a.output)
print(a.output)
print('Chapter pages:',page_map)
